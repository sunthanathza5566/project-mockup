#!/usr/bin/env python3
"""Build the research report (เล่มวิจัย) from Markdown chapters.

Usage:
  build.py make-ref   create thesis/templates/reference.docx (TH Sarabun New, Thai margins, sizes from STYLE_SIZES)
  build.py diagrams [--force]  render docs/diagrams/*.mmd -> thesis/export/figures/*.png (needs mmdc;
                      --force = render ใหม่ทุกไฟล์ ไม่ดู mtime)
  build.py docx       thesis/export/thesis.docx (needs pandoc)
  build.py gemini     thesis/export/GEMINI_HANDOFF.md (no dependencies)
  build.py check      report missing chapters, open markers (by type), missing diagrams/images
  build.py verify     open thesis.docx with python-docx: chapter headings, figure count, fonts
  build.py all [--force]  make-ref + diagrams + docx + gemini + check + verify

pandoc: uses `pandoc` from PATH, else the binary bundled with pypandoc_binary
  (pip install --user pypandoc_binary)

Diagram placeholder inside a chapter:
  {{diagram:seq-login|รูปที่ 3.1 Sequence diagram การเข้าสู่ระบบ}}
  {{diagram:er-aunqa|รูปที่ 3.13 ...|landscape}}   วางรูปในหน้าแนวนอน (section แยก) ให้ตัวอักษรใหญ่ขึ้น
"""
import re
import shutil
import subprocess
import sys
import warnings
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parents[2]
TH = ROOT / "thesis"
CH = TH / "chapters"
OUT = TH / "export"
FIG = OUT / "figures"
DIAG = ROOT / "docs" / "diagrams"
META = TH / "meta.md"
REF = TH / "templates" / "reference.docx"
PH = re.compile(r"\{\{diagram:([A-Za-z0-9_\-]+)(?:\|([^}|]*))?(?:\|(landscape))?\}\}")
FIG_NO_RE = re.compile(r"^(รูปที่\s*[0-9ก-ฮ]+\.[0-9]+)")
# ป้ายค้าง: นับทั้งแบบเปล่า [ต้องการข้อมูล] และแบบมีรายละเอียด [ต้องการข้อมูล: ...]
MARKER_RE = re.compile(
    r"\[(TODO|ต้องตรวจสอบกับคู่มือสาขา|ต้องตรวจสอบ|ต้องการข้อมูล|ต้องการอ้างอิง)(?:[:：][^\]]*)?\]")
IMG_RE = re.compile(r"!\[[^\]]*\]\(([^)\s]+)\)")
COMMENT_RE = re.compile(r"<!--.*?-->", re.S)
CHAPTER_HEADINGS = ["บทที่ 1", "บทที่ 2", "บทที่ 3", "บทที่ 4", "บทที่ 5", "บรรณานุกรม", "ภาคผนวก"]

EXPECTED = [
    "00-front-matter.md", "01-introduction.md", "02-literature.md",
    "03-methodology.md", "04-results.md", "05-conclusion.md",
    "90-references.md", "95-appendix.md",
]

THAI_FONT = "TH Sarabun New"
# ค่ากลางจาก .agents/research/standards/thesis-th.md หัวข้อ 2.2 (ชื่อบท 18 หนา, 1.1 = 16 หนา, 1.1.1 = 16 ธรรมดา)
STYLE_SIZES = {
    "Normal": 16, "Body Text": 16, "First Paragraph": 16, "Compact": 16,
    "Block Text": 16, "Caption": 16, "Image Caption": 16, "Table Caption": 16,
    "Title": 20, "Subtitle": 18, "Heading 1": 18, "Heading 2": 16,
    "Heading 3": 16, "Heading 4": 16, "TOC Heading": 18,
}
BOLD_STYLES = {"Title", "Subtitle", "Heading 1", "Heading 2", "TOC Heading"}  # 1.1.1 ลงไป = ธรรมดา
BODY_STYLES = ("Body Text", "First Paragraph")  # กระจายไทย + ย่อหน้า 0.5"

GEMINI_HEADER = """# คำสั่งสำหรับ Gemini — จัดทำเล่มวิจัยจากเนื้อหาด้านล่าง

คุณคือผู้ช่วยจัดรูปแบบเอกสารวิชาการภาษาไทย หน้าที่ของคุณคือ **นำเนื้อหาด้านล่างไปจัดทำเป็นเล่มวิจัยใน Google Docs**
โดยไม่เปลี่ยนสาระสำคัญ

## รูปแบบที่ต้องใช้
- ฟอนต์ TH Sarabun New ขนาด 16 pt (ชื่อบท 18 pt ตัวหนา จัดกึ่งกลาง, หัวข้อ 1.1 = 16 pt ตัวหนา, หัวข้อ 1.1.1 = 16 pt ธรรมดา) [ต้องตรวจสอบกับคู่มือสาขา]
- ระยะขอบ: บน 1.5 นิ้ว, ซ้าย 1.5 นิ้ว, ล่าง 1 นิ้ว, ขวา 1 นิ้ว / ระยะบรรทัด 1 เท่า
- จัดข้อความแบบกระจายไทย ย่อหน้าบรรทัดแรก 0.5 นิ้ว
- ขึ้นบทใหม่ (บทที่ 1–5) ที่หน้าใหม่เสมอ
- ลำดับส่วนนำ: ปกนอก → ปกใน → ใบรับรอง → บทคัดย่อ → Abstract → กิตติกรรมประกาศ → สารบัญ → สารบัญตาราง → สารบัญรูป → คำอธิบายสัญลักษณ์และคำย่อ; เลขหน้าส่วนนำ ก ข ค / เนื้อหาเริ่ม 1 ที่บทที่ 1 มุมบนขวา
- เลขหัวข้อตามที่ให้มา (1.1, 1.2.1 ...) / เลขรูปและตารางตามบท (รูปที่ 3.1, ตารางที่ 4.2)
- คำบรรยายรูปอยู่ใต้รูป คำบรรยายตารางอยู่เหนือตาราง (บรรทัด `Table: ...` คือคำบรรยายตาราง)
- สร้างสารบัญ สารบัญตาราง และสารบัญรูป
- บรรณานุกรมตามรูปแบบ APA 7th

## Diagram
บล็อก ```mermaid``` คือโค้ด diagram ให้แปลงเป็นรูป (เช่นผ่าน mermaid.live) แล้ววางแทนบล็อก พร้อมคำบรรยายรูปที่อยู่ด้านบนบล็อก (ย้ายคำบรรยายไปไว้ใต้รูป)
- ห้ามลบบรรทัด `%%{init: …}%%` ที่มี themeCSS `.edge-pattern-dashed` (แก้บั๊กเส้นประของ mermaid 12.2.0 — ถ้าลบ เส้น include/extend/dependency จะกลายเป็นเส้นทึบ)
- รูปที่มีหมายเหตุ "(จัดวางในหน้าแนวนอน)" ให้วางรูปและคำบรรยายในหน้าแนวนอนแยก [ต้องตรวจสอบกับคู่มือสาขา]
- รูปภาพหน้าจอ `![...](../../.agents/screenshots/...)` ใช้ไฟล์ภาพจากโฟลเดอร์นั้นในโปรเจกต์ (thesis.docx มีภาพฝังอยู่แล้ว)

## ข้อห้าม
- ห้ามแต่งข้อมูล ตัวเลข ผลการทดลอง หรือแหล่งอ้างอิงขึ้นใหม่
- ข้อความที่มีป้าย [ต้องตรวจสอบ...] [ต้องตรวจสอบกับคู่มือสาขา...] [ต้องการข้อมูล...] [ต้องการอ้างอิง] [TODO...] ให้คงไว้และไฮไลต์สีเหลือง
- ข้อความใน comment `<!-- source: ... -->` คือที่มาของตัวเลข ไม่ต้องแสดงในเล่ม แต่ห้ามเปลี่ยนตัวเลข
- ปรับภาษาได้เฉพาะให้เป็นภาษาวิชาการที่ถูกต้อง สละสลวย โดยความหมายเดิม

## เมื่อเสร็จ
สรุปสั้น ๆ ว่าปรับภาษาจุดใด และมีป้ายที่ยังค้างกี่จุด

---
"""


def pandoc_bin():
    exe = shutil.which("pandoc")
    if exe:
        return exe
    try:
        import pypandoc  # pypandoc_binary ติดตั้งระดับ user
        return pypandoc.get_pandoc_path()
    except Exception:
        return None


def chapters():
    return sorted(CH.glob("*.md")) if CH.exists() else []


def mmd_path(name):
    return DIAG / f"{name}.mmd"


def assemble(mode):
    parts = []
    if META.exists():
        parts.append(META.read_text(encoding="utf-8"))
    for p in chapters():
        parts.append(p.read_text(encoding="utf-8"))
    text = "\n\n".join(parts)

    def repl(m):
        name, caption = m.group(1), (m.group(2) or m.group(1)).strip()
        png = FIG / f"{name}.png"
        src = mmd_path(name)
        if mode == "docx" and png.exists():
            return f"![{caption}](figures/{name}.png)"
        code = src.read_text(encoding="utf-8").strip() if src.exists() else "%% [ต้องการข้อมูล] ไม่พบไฟล์ diagram"
        note = " (จัดวางในหน้าแนวนอน)" if m.group(3) else ""
        return f"**{caption}**{note}\n\n```mermaid\n{code}\n```"

    return PH.sub(repl, text)


def _set_fonts(rpr, OxmlElement, qn):
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for attr in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rf.set(qn(attr), THAI_FONT)
    for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        rf.attrib.pop(qn(attr), None)


def cmd_make_ref():
    pandoc = pandoc_bin()
    if not pandoc:
        sys.exit("ต้องติดตั้ง pandoc ก่อน (หรือ pip install --user pypandoc_binary)")
    try:
        from docx import Document
        from docx.enum.text import WD_ALIGN_PARAGRAPH
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches, Pt
    except ImportError:
        sys.exit("ต้องติดตั้ง python-docx ก่อน: pip install python-docx")
    REF.parent.mkdir(parents=True, exist_ok=True)
    subprocess.run([pandoc, "-o", str(REF), "--print-default-data-file", "reference.docx"],
                   check=True)
    d = Document(str(REF))
    warnings.filterwarnings("ignore", category=UserWarning)
    for name, size in STYLE_SIZES.items():
        try:
            st = d.styles[name]
        except KeyError:
            continue
        st.font.size = Pt(size)
        rpr = st.element.get_or_add_rPr()
        _set_fonts(rpr, OxmlElement, qn)
        szcs = rpr.find(qn("w:szCs"))
        if szcs is None:
            szcs = OxmlElement("w:szCs")
            rpr.append(szcs)
        szcs.set(qn("w:val"), str(size * 2))
        if name.startswith("Heading") or name in BOLD_STYLES:
            st.font.bold = name in BOLD_STYLES
            st.font.italic = False
            st.font.color.rgb = None
        pf = st.paragraph_format
        pf.line_spacing = 1.0
        if name in ("Heading 1", "TOC Heading", "Title", "Subtitle"):
            pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if name == "Heading 1":
            pf.page_break_before = True
        if name in ("Image Caption", "Caption"):
            pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if name == "Table Caption":
            pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
        if name in BODY_STYLES:
            pf.first_line_indent = Inches(0.5)
            ppr = st.element.get_or_add_pPr()
            jc = ppr.find(qn("w:jc"))
            if jc is None:
                jc = OxmlElement("w:jc")
                ppr.append(jc)
            jc.set(qn("w:val"), "thaiDistribute")
    # ค่าเริ่มต้นของทั้งเอกสาร (ตาราง/ข้อความที่ไม่มีสไตล์) = TH Sarabun New
    defaults = d.styles.element.find(qn("w:docDefaults"))
    if defaults is not None:
        rpr = defaults.find(qn("w:rPrDefault") + "/" + qn("w:rPr"))
        if rpr is not None:
            _set_fonts(rpr, OxmlElement, qn)
    for s in d.sections:
        s.page_width = Inches(8.27)   # A4
        s.page_height = Inches(11.69)
        s.top_margin = Inches(1.5)
        s.left_margin = Inches(1.5)
        s.bottom_margin = Inches(1)
        s.right_margin = Inches(1)
    d.save(str(REF))
    print(f"wrote {REF.relative_to(ROOT)}")


def cmd_diagrams(force=False):
    mmdc = shutil.which("mmdc")
    if not mmdc:
        print("ข้าม: ไม่พบ mmdc (ติดตั้ง: npm i -g @mermaid-js/mermaid-cli) — docx จะใส่โค้ด mermaid แทนรูป")
        return
    FIG.mkdir(parents=True, exist_ok=True)
    for src in sorted(DIAG.glob("*.mmd")):
        out = FIG / f"{src.stem}.png"
        if not force and out.exists() and out.stat().st_mtime >= src.stat().st_mtime:
            continue
        r = subprocess.run([mmdc, "-i", str(src), "-o", str(out), "-s", "2", "-b", "white"],
                           capture_output=True, text=True, encoding="utf-8", errors="replace")
        print(("rendered " if r.returncode == 0 else "FAILED ") + src.name)
        if r.returncode != 0:
            print(r.stderr[-800:])


def cmd_docx():
    pandoc = pandoc_bin()
    if not pandoc:
        sys.exit("ต้องติดตั้ง pandoc ก่อน (apt install pandoc / brew install pandoc / pip install --user pypandoc_binary)")
    OUT.mkdir(parents=True, exist_ok=True)
    tmp = OUT / "_assembled.md"
    tmp.write_text(assemble("docx"), encoding="utf-8")
    sep = ";" if sys.platform == "win32" else ":"
    args = [pandoc, str(tmp), "-o", str(OUT / "thesis.docx"),
            "--toc", "--toc-depth=2", "-M", "toc-title=สารบัญ",
            "--lof", "--lot", "-M", "lof-title=สารบัญรูป", "-M", "lot-title=สารบัญตาราง",
            f"--resource-path={OUT}{sep}{CH}"]
    if REF.exists():
        args.append(f"--reference-doc={REF}")
    else:
        print("เตือน: ยังไม่มี reference.docx — รัน build.py make-ref เพื่อให้ได้ฟอนต์ TH Sarabun New")
    subprocess.run(args, check=True)
    tmp.unlink()
    fit_images(OUT / "thesis.docx")
    print(f"wrote {(OUT / 'thesis.docx').relative_to(ROOT)}")


MAX_IMG_HEIGHT_IN = 8.4  # พื้นที่พิมพ์ A4 สูง 9.19" (ขอบ 1.5/1) เผื่อคำบรรยายใต้รูป
# หน้าแนวนอน: A4 หมุน 11.69" x 8.27" ใช้ขอบเดิม (บน 1.5 ซ้าย 1.5 ล่าง 1 ขวา 1)
# -> พื้นที่พิมพ์ 9.19" x 5.77" เผื่อคำบรรยายใต้รูป [ต้องตรวจสอบกับคู่มือสาขา: ขอบของหน้าแนวนอน]
LAND_MAX_W_IN = 9.19
LAND_MAX_H_IN = 5.2


def landscape_figures():
    """เลขรูป (เช่น 'รูปที่ 3.13') ที่ placeholder ระบุ |landscape"""
    out = set()
    for p in chapters():
        t = COMMENT_RE.sub("", p.read_text(encoding="utf-8"))
        for m in PH.finditer(t):
            if m.group(3):
                n = FIG_NO_RE.match((m.group(2) or "").strip())
                if n:
                    out.add(re.sub(r"\s+", " ", n.group(1)))
    return out


def _para_text(p, qn):
    return "".join(t.text or "" for t in p.iter(qn("w:t"))).strip()


def fit_images(path):
    """ย่อรูปที่สูงเกินหน้ากระดาษ (pandoc ย่อเฉพาะความกว้าง) และวางรูปที่ระบุ landscape ในหน้าแนวนอน"""
    try:
        import copy
        from docx import Document
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.shared import Inches
    except ImportError:
        print("เตือน: ไม่มี python-docx — ข้ามการย่อรูปที่สูงเกินหน้า")
        return
    d = Document(str(path))
    body = d.element.body
    body_sect = body.find(qn("w:sectPr"))
    want_land = landscape_figures()
    els = list(body)
    changed, land_done = 0, []

    def extents(p):
        return [e for e in p.iter(qn("wp:extent"))] + [e for e in p.iter(qn("a:ext"))]

    def set_size(p, cx, cy):
        for e in extents(p):
            e.set("cx", str(int(cx)))
            e.set("cy", str(int(cy)))

    def add_sect(p, landscape):
        ppr = p.find(qn("w:pPr"))
        if ppr is None:
            ppr = OxmlElement("w:pPr")
            p.insert(0, ppr)
        if ppr.find(qn("w:sectPr")) is not None:
            return
        s = copy.deepcopy(body_sect)
        sz = s.find(qn("w:pgSz"))
        w, h = int(sz.get(qn("w:w"))), int(sz.get(qn("w:h")))
        if landscape:
            sz.set(qn("w:w"), str(max(w, h)))
            sz.set(qn("w:h"), str(min(w, h)))
            sz.set(qn("w:orient"), "landscape")
        ppr.append(s)

    for i, el in enumerate(els):
        if el.tag != qn("w:p") or not el.findall(".//" + qn("a:blip")):
            continue
        ext = el.find(".//" + qn("wp:extent"))
        if ext is None:
            continue
        cx, cy = int(ext.get("cx")), int(ext.get("cy"))
        cap = els[i + 1] if i + 1 < len(els) else None
        cap_txt = _para_text(cap, qn) if cap is not None else ""
        n = FIG_NO_RE.match(cap_txt)
        fig_no = re.sub(r"\s+", " ", n.group(1)) if n else None
        if fig_no in want_land and cap is not None and i > 0:
            ratio = min(Inches(LAND_MAX_W_IN) / cx, Inches(LAND_MAX_H_IN) / cy)
            set_size(el, cx * ratio, cy * ratio)
            prev = els[i - 1]
            if prev.tag == qn("w:p"):
                add_sect(prev, landscape=False)   # ปิด section แนวตั้งก่อนรูป
            add_sect(cap, landscape=True)         # section แนวนอน = รูป + คำบรรยาย
            land_done.append(fig_no)
            changed += 1
        elif cy > Inches(MAX_IMG_HEIGHT_IN):
            ratio = Inches(MAX_IMG_HEIGHT_IN) / cy
            set_size(el, cx * ratio, cy * ratio)
            changed += 1
    missing = sorted(want_land - set(land_done))
    if missing:
        print("เตือน: ไม่พบรูปที่ต้องวางแนวนอน: " + ", ".join(missing))
    if changed:
        d.save(str(path))
        print(f"ปรับขนาดรูป {changed} รูป (หน้าแนวนอน {len(land_done)}: {', '.join(land_done) or '-'})")


def cmd_gemini():
    OUT.mkdir(parents=True, exist_ok=True)
    body = assemble("gemini")
    stamp = f"_สร้างเมื่อ {date.today().isoformat()} จากไฟล์ใน thesis/chapters/_\n\n"
    path = OUT / "GEMINI_HANDOFF.md"
    path.write_text(GEMINI_HEADER + stamp + body + "\n", encoding="utf-8")
    print(f"wrote {path.relative_to(ROOT)}")


def cmd_check():
    have = {p.name for p in chapters()}
    missing = [n for n in EXPECTED if n not in have]
    print("บทที่ยังไม่มีไฟล์:", ", ".join(missing) if missing else "ครบ")
    total, problems, placeholders, images = 0, 0, 0, 0
    by_type = {}
    seen = {}
    for p in ([META] if META.exists() else []) + chapters():
        t = COMMENT_RE.sub("", p.read_text(encoding="utf-8"))
        counts = {}
        for m in MARKER_RE.finditer(t):
            k = "[" + m.group(1) + "]"
            counts[k] = counts.get(k, 0) + 1
            by_type[k] = by_type.get(k, 0) + 1
        for m in PH.finditer(t):
            placeholders += 1
            name = m.group(1)
            if name in seen:
                print(f"  เตือน {p.name}: diagram {name} ใช้ซ้ำ (ครั้งแรกใน {seen[name]})")
            seen.setdefault(name, p.name)
            if not mmd_path(name).exists():
                problems += 1
                print(f"  ERROR {p.name}: ไม่พบไฟล์ .mmd ของ diagram: {name}")
            elif not (FIG / f"{name}.png").exists():
                problems += 1
                print(f"  ERROR {p.name}: ยังไม่ได้ render PNG: {name} (รัน build.py diagrams)")
        for m in IMG_RE.finditer(t):
            images += 1
            if not (CH / m.group(1)).resolve().exists():
                problems += 1
                print(f"  ERROR {p.name}: ไม่พบไฟล์รูป {m.group(1)}")
        if counts:
            total += sum(counts.values())
            print(f"  {p.name}: " + ", ".join(f"{k} x{v}" for k, v in counts.items()))
    unused = sorted(s.stem for s in DIAG.glob("*.mmd") if s.stem not in seen)
    if unused:
        print("  diagram ที่ยังไม่ได้ใช้ในเล่ม: " + ", ".join(unused))
    print("ป้ายแยกชนิด: " + (", ".join(f"{k} {v}" for k, v in sorted(by_type.items())) or "-"))
    print(f"ป้ายค้างทั้งหมด: {total}")
    print(f"diagram placeholder: {placeholders} · รูปจาก Markdown: {images} · ปัญหารูป/diagram: {problems}")
    return problems


def cmd_verify():
    """เปิด thesis.docx ตรวจหัวบท จำนวนรูป และฟอนต์"""
    try:
        from docx import Document
        from docx.oxml.ns import qn
    except ImportError:
        sys.exit("ต้องติดตั้ง python-docx ก่อน: pip install python-docx")
    path = OUT / "thesis.docx"
    if not path.exists():
        sys.exit("ไม่พบ thesis/export/thesis.docx — รัน build.py docx ก่อน")
    d = Document(str(path))
    ok = True
    h1 = [p.text.strip() for p in d.paragraphs if p.style.name == "Heading 1"]
    for want in CHAPTER_HEADINGS:
        hit = any(h.startswith(want) for h in h1)
        ok &= hit
        print(f"  หัวบท {want}: {'มี' if hit else 'ไม่พบ'}")
    text = COMMENT_RE.sub("", "\n\n".join(p.read_text(encoding="utf-8") for p in chapters()))
    expected = len(PH.findall(text)) + len(IMG_RE.findall(text))
    blips = len(d.element.body.findall(".//" + qn("a:blip")))
    print(f"  รูปใน docx: {blips} · placeholder + รูป Markdown ในต้นฉบับ: {expected}")
    ok &= blips == expected
    land = [s for s in d.element.body.iter(qn("w:pgSz")) if s.get(qn("w:orient")) == "landscape"]
    want = len(landscape_figures())
    print(f"  section หน้าแนวนอน: {len(land)} · placeholder |landscape: {want}")
    ok &= len(land) == want
    tall = [round(int(e.get("cy")) / 914400, 2) for e in d.element.body.iter(qn("wp:extent"))
            if int(e.get("cy")) > 914400 * MAX_IMG_HEIGHT_IN + 1]
    print(f"  รูปที่สูงเกิน {MAX_IMG_HEIGHT_IN}\": {len(tall)}")
    ok &= not tall
    for name in ("Normal", "Heading 1", "Heading 2", "Heading 3", "Body Text",
                 "First Paragraph", "Compact", "Image Caption", "Table Caption"):
        try:
            st = d.styles[name]
        except KeyError:
            continue
        rf = st.element.find(qn("w:rPr") + "/" + qn("w:rFonts"))
        cs = rf.get(qn("w:cs")) if rf is not None else None
        ok &= cs == THAI_FONT
        size = st.font.size.pt if st.font.size else None
        print(f"  สไตล์ {name}: ฟอนต์ {cs} {size}pt หนา={st.font.bold}")
    run_fonts = set()
    for el in d.element.body.iter(qn("w:rFonts")):
        for attr in ("w:ascii", "w:cs"):
            if el.get(qn(attr)):
                run_fonts.add(el.get(qn(attr)))
    if run_fonts - {THAI_FONT}:
        print(f"  หมายเหตุ: ข้อความบางส่วนกำหนดฟอนต์เอง: {sorted(run_fonts - {THAI_FONT})}")
    print("verify: " + ("ผ่าน" if ok else "ไม่ผ่าน"))
    return ok


def main(argv):
    cmds = {"make-ref": cmd_make_ref, "diagrams": cmd_diagrams, "docx": cmd_docx,
            "gemini": cmd_gemini, "check": cmd_check, "verify": cmd_verify}
    if not argv:
        sys.exit(__doc__)
    force = "--force" in argv[1:]
    if argv[0] == "all":
        cmd_make_ref(); cmd_diagrams(force); cmd_docx(); cmd_gemini()
        problems = cmd_check()
        ok = cmd_verify()
        if problems or not ok:
            sys.exit(1)
    elif argv[0] == "diagrams":
        cmd_diagrams(force)
    elif argv[0] in cmds:
        cmds[argv[0]]()
    else:
        sys.exit(__doc__)


if __name__ == "__main__":
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8")
    main(sys.argv[1:])
